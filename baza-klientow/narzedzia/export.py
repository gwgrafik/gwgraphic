import sys, json, csv, collections, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
OUT = sys.argv[1]
rows = json.load(open(f'{OUT}/_rows.json'))
COLS = ['Priorytet','Punkty','Firma','Miasto','Adres','Strona WWW','Stan strony','HTTPS','Na telefon (viewport)','Rok w stopce','System strony',
        'Social media (znalezione)','Oklejenie aut / odzież / witryna','Sygnały zakupu','Pracownicy (LinkedIn)','Zmiana zatrudnienia r/r','Rok założenia',
        'Sieć / franczyza','Ocena ogólna','Luki','Szansa dla GW','Status kontaktu','Źródło','Sprawdzono']
ORDER = ["Dekarze","Budowlanka i remonty","Malarze","Instalatorzy i elektrycy","Ogrody i zieleń","Sprzątanie","Transport i logistyka","Motoryzacja","Szkoły jazdy",
         "Inne rzemiosło","Fryzjerzy i beauty","Zdrowie i fizjoterapia","Fitness i sport","Gastronomia","Sklepy lokalne","Biura i usługi B2B","Partnerzy (dostawcy)"]
cats = [c for c in ORDER if any(r['Kategoria']==c for r in rows)] + sorted({r['Kategoria'] for r in rows} - set(ORDER))
key = lambda r: ({'A':0,'B':1,'C':2}[r['Priorytet']], -r['Punkty'], r['Firma'].lower())
def slug(s): return re.sub(r'[^a-z0-9]+','-', s.lower().translate(str.maketrans('ąćęłńóśźż','acelnoszz'))).strip('-')
wb = Workbook(); ws0 = wb.active; ws0.title = 'Podsumowanie'
H = Font(bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='222222')
PF = {'A': PatternFill('solid', fgColor='C6EFCE'), 'B': PatternFill('solid', fgColor='FFEB9C'), 'C': PatternFill('solid', fgColor='EEEEEE')}
summ = [['Kategoria','Firm','Priorytet A','Priorytet B','Strona OK','Strona przestarzała','Strona z błędem','Brak strony (dane firmy)','Nie znaleziono strony w danych','Nieustalone','Z sygnałem zakupu','Sieci']]
for c in cats:
    rs = [r for r in rows if r['Kategoria']==c]
    summ.append([c, len(rs), sum(r['Priorytet']=='A' for r in rs), sum(r['Priorytet']=='B' for r in rs),
        sum(r['Stan strony']=='OK' for r in rs), sum(r['Stan strony']=='PRZESTARZAŁA' for r in rs), sum(r['Stan strony']=='BŁĄD STRONY (do potwierdzenia)' for r in rs),
        sum(r['Stan strony']=='BRAK STRONY' for r in rs), sum(r['Stan strony'].startswith('NIE ZNALEZIONO') for r in rs), sum(r['Stan strony'].startswith('NIEUSTALONE') for r in rs), sum(bool(r['Sygnały zakupu']) for r in rs), sum(bool(r['Sieć / franczyza']) for r in rs)])
summ.append(['RAZEM', len(rows)] + [sum(x[i] for x in summ[1:]) for i in range(2, 12)])
for row in summ: ws0.append(row)
for cell in ws0[1]: cell.font = H; cell.fill = HF
ws0.column_dimensions['A'].width = 28
for i in range(2, 13): ws0.column_dimensions[get_column_letter(i)].width = 16
ws0.append([]); ws0.append(['Legenda: Priorytet A = najlepsze dopasowanie + luki/sygnały; „Oklejenie aut / odzież / witryna” zawsze NIEUSTALONE bez zdjęć lub wizji lokalnej; „nie znaleziono social media” ≠ brak.'])
def sheet(title, rs):
    ws = wb.create_sheet(title[:31]); ws.append(COLS)
    for cell in ws[1]: cell.font = H; cell.fill = HF; cell.alignment = Alignment(wrap_text=True, vertical='top')
    for r in sorted(rs, key=key):
        ws.append([r.get(k, '') for k in COLS]); ws.cell(ws.max_row, 1).fill = PF[r['Priorytet']]
    widths = {'Firma':34,'Miasto':16,'Adres':24,'Strona WWW':32,'Stan strony':18,'Social media (znalezione)':22,'Oklejenie aut / odzież / witryna':22,'Sygnały zakupu':34,'Luki':40,'Szansa dla GW':48,'Źródło':30,'Ocena ogólna':18}
    for i, k in enumerate(COLS, 1): ws.column_dimensions[get_column_letter(i)].width = widths.get(k, 11)
    ws.freeze_panes = 'D2'; ws.auto_filter.ref = ws.dimensions
cd = json.load(open(f'{OUT}/_cand.json'))
wk = wb.create_sheet('KANDYDACI – wszystkie branże', 1); wk.append(cd['hdr'])
for cell in wk[1]: cell.font = H; cell.fill = HF; cell.alignment = Alignment(wrap_text=True, vertical='top')
for r in cd['rows']: wk.append(r)
for i, w in enumerate([22,34,16,34,38,22,34,34,50,12,44,12,12,30,9], 1): wk.column_dimensions[get_column_letter(i)].width = w
wk.freeze_panes = 'C2'; wk.auto_filter.ref = wk.dimensions
top = [r for r in rows if r['Priorytet']=='A']
sheet('TOP – priorytet A', top)
for c in cats: sheet(c, [r for r in rows if r['Kategoria']==c])
wb.save(f'{OUT}/baza-klientow.xlsx')
import os; os.makedirs(f'{OUT}/csv', exist_ok=True)
for c in cats:
    with open(f'{OUT}/csv/{slug(c)}.csv', 'w', encoding='utf-8-sig', newline='') as f:
        w = csv.writer(f, delimiter=';'); w.writerow(['Kategoria'] + COLS)
        for r in sorted([r for r in rows if r['Kategoria']==c], key=key): w.writerow([c] + [r.get(k, '') for k in COLS])
with open(f'{OUT}/csv/KANDYDACI-wszystkie-branze.csv', 'w', encoding='utf-8-sig', newline='') as f:
    w = csv.writer(f, delimiter=';'); w.writerow(cd['hdr']); w.writerows(cd['rows'])
json.dump(summ, open(f'{OUT}/_summary.json','w'), ensure_ascii=False)
print('ok', len(rows), len(top))
